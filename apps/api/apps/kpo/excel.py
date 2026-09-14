"""
KPO Excel export (master prompt §20). openpyxl is the only Excel-writing
dependency in the project — added specifically for this, not speculative.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

HEADERS = [
    "Redni broj",
    "Datum",
    "Opis",
    "Prihod od prodaje proizvoda",
    "Prihod od izvršenih usluga",
    "Svega prihodi",
]


def render_kpo_excel(company, entries) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "KPO"

    bold = Font(bold=True)

    header_fields = [
        ("Firma-radnja", company.name),
        ("PIB", company.pib or ""),
        ("Obveznik", company.legal_name or company.name),
        ("Sedište", f"{company.address}, {company.city}" if company.address else ""),
        ("Šifra delatnosti", company.activity_code or ""),
    ]
    for label, value in header_fields:
        sheet.append([label, value])
        sheet.cell(row=sheet.max_row, column=1).font = bold
    sheet.append([])

    sheet.append(HEADERS)
    for cell in sheet[sheet.max_row]:
        cell.font = bold

    for entry in entries:
        sheet.append(
            [
                entry.sequence_number,
                entry.entry_date,
                entry.description,
                float(entry.product_revenue),
                float(entry.service_revenue),
                float(entry.total_revenue),
            ]
        )

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
