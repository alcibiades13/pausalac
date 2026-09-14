from io import BytesIO

from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook

from apps.clients.models import Client
from apps.companies.models import Company
from apps.invoices.services import InvoiceService
from apps.kpo.excel import render_kpo_excel
from apps.kpo.models import KPOEntry


class RenderKpoExcelTests(TestCase):
    def test_header_and_row_values(self):
        company = Company.objects.create(
            name="Firma A",
            pib="123456789",
            legal_name="Firma A PR",
            address="Ulica 1",
            city="Beograd",
        )
        client_obj = Client.objects.create(company=company, name="Klijent A")
        today = timezone.now().date()
        invoice = InvoiceService.create_draft(
            company=company,
            client=client_obj,
            issue_date=today,
            due_date=today,
            service_date=today,
            notes="",
            line_items=[
                {
                    "description": "Widget",
                    "item_type": "product",
                    "quantity": "1",
                    "unit_price": "300.00",
                },
                {
                    "description": "Konsalting",
                    "item_type": "service",
                    "quantity": "1",
                    "unit_price": "700.00",
                },
            ],
        )
        InvoiceService.issue(invoice, actor=None)
        entries = KPOEntry.objects.for_company(company)

        xlsx_bytes = render_kpo_excel(company, entries)
        workbook = load_workbook(BytesIO(xlsx_bytes))
        sheet = workbook.active

        values = [[cell.value for cell in row] for row in sheet.iter_rows()]
        flattened_text = "\n".join(str(v) for row in values for v in row if v is not None)
        self.assertIn("123456789", flattened_text)
        self.assertIn("Firma A", flattened_text)
        self.assertIn("Redni broj", flattened_text)
        self.assertIn("Prihod od prodaje proizvoda", flattened_text)

        data_row = next(row for row in values if row[0] == 1)
        self.assertEqual(data_row[3], 300.0)
        self.assertEqual(data_row[4], 700.0)
        self.assertEqual(data_row[5], 1000.0)
